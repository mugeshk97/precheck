"""Batch-run all FA PDFs through the pipeline and save results to Excel."""

import asyncio
import glob
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pipeline import run_pipeline


def save_to_excel(results: list[dict], output_path: str) -> str:
    wb = Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    match_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    no_match_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    no_isi_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    headers = ["Final Asset", "ISI Selected", "Drug", "Audience", "Coverage", "Authenticity", "F1", "Match Category", "Sections"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for row_idx, r in enumerate(results, 2):
        if "error" in r:
            ws.cell(row=row_idx, column=1, value=r["fa"])
            ws.cell(row=row_idx, column=2, value=f"ERROR: {r['error']}")
            continue

        o = r.get("overall", {})
        ws.cell(row=row_idx, column=1, value=r.get("fa", ""))
        ws.cell(row=row_idx, column=2, value=r.get("isi") or "None")
        ws.cell(row=row_idx, column=3, value=r.get("drug", ""))
        ws.cell(row=row_idx, column=4, value=r.get("audience", ""))
        ws.cell(row=row_idx, column=5, value=o.get("coverage", 0))
        ws.cell(row=row_idx, column=6, value=o.get("authenticity", 0))
        ws.cell(row=row_idx, column=7, value=o.get("f1", 0))

        match_cat = r.get("match_category", "")
        match_cell = ws.cell(row=row_idx, column=8, value=match_cat)
        if match_cat == "Closest Match":
            match_cell.fill = match_fill
        elif match_cat == "No ISI":
            match_cell.fill = no_isi_fill
        else:
            match_cell.fill = no_match_fill

        ws.cell(row=row_idx, column=9, value=len(r.get("sections", [])))

        for col in [5, 6, 7]:
            ws.cell(row=row_idx, column=col).number_format = "0.0"

    # Auto-width
    for col_idx in range(1, len(headers) + 1):
        max_len = len(headers[col_idx - 1])
        for row_idx in range(2, len(results) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(results) + 1}"

    # ── Sheet 2: Section Details ──────────────────────────────────────────
    ws2 = wb.create_sheet("Section Details")
    sec_headers = ["Final Asset", "Section", "Coverage", "Authenticity", "F1", "ISI Sentences", "FA Fragments", "Mismatches"]
    for col, h in enumerate(sec_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row = 2
    for r in results:
        if "error" in r or not r.get("sections"):
            continue
        fa_name = r.get("fa", "")
        for sec in r["sections"]:
            ws2.cell(row=row, column=1, value=fa_name)
            ws2.cell(row=row, column=2, value=sec.get("title", ""))
            ws2.cell(row=row, column=3, value=sec.get("coverage", 0))
            ws2.cell(row=row, column=4, value=sec.get("authenticity", 0))
            ws2.cell(row=row, column=5, value=sec.get("f1", 0))
            ws2.cell(row=row, column=6, value=sec.get("isi_sentence_count", 0))
            ws2.cell(row=row, column=7, value=sec.get("fa_fragment_count", 0))
            ws2.cell(row=row, column=8, value=len(sec.get("mismatches", [])))
            for col in [3, 4, 5]:
                ws2.cell(row=row, column=col).number_format = "0.0"
            row += 1

    for col_idx in range(1, len(sec_headers) + 1):
        max_len = len(sec_headers[col_idx - 1])
        for r in range(2, row):
            val = ws2.cell(row=r, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws2.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

    ws2.auto_filter.ref = f"A1:{get_column_letter(len(sec_headers))}{row - 1}"

    wb.save(output_path)
    return output_path


async def main():
    fa_files = sorted(glob.glob("finalassets/*.pdf"))
    print(f"Found {len(fa_files)} FA files\n")

    results = []
    for fa_path in fa_files:
        print(f"\n{'='*70}")
        print(f"Processing: {Path(fa_path).name}")
        print(f"{'='*70}")
        try:
            report = await run_pipeline(fa_path=fa_path, isi_dir="isi", debug=True)
            results.append(report)
        except Exception as e:
            print(f"  ERROR: {e}")
            results.append({"fa": Path(fa_path).name, "error": str(e)})

    # Save to Excel
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    excel_path = save_to_excel(results, f"pipeline_results_{timestamp}.xlsx")
    print(f"\n{'='*70}")
    print(f"Results saved to: {excel_path}")

    # Console summary
    print(f"\n{'FA':<65} {'ISI':<45} {'Cov':>5} {'Auth':>5} {'F1':>5}  {'Match'}")
    print(f"{'-'*130}")
    for r in results:
        if "error" in r:
            print(f"{r['fa']:<65} ERROR: {r['error']}")
            continue
        o = r.get("overall", {})
        isi_name = r.get("isi") or "None"
        print(
            f"{r['fa']:<65} {isi_name:<45} "
            f"{o.get('coverage',0):5.1f} {o.get('authenticity',0):5.1f} {o.get('f1',0):5.1f}  "
            f"{r.get('match_category','?')}"
        )


if __name__ == "__main__":
    asyncio.run(main())
